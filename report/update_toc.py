"""Populate the cache of a live Word TOC field from headings and rendered page numbers.
No manually typed contents or page numbers. Word can refresh the TOC normally.
"""
from pathlib import Path
import re
import sys
import pdfplumber
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

root=Path(__file__).resolve().parents[1]
path=root/'Отчёт ДЗ1 — Невокшенов.docx'
pdf=Path(sys.argv[1])
doc=Document(path)
normalize=lambda s:re.sub(r'\s+',' ',s).strip()
with pdfplumber.open(pdf) as rendered:
    pages=[normalize(p.extract_text() or '') for p in rendered.pages]
entries=[]
for p in doc.paragraphs:
    if p.style.name not in ('Heading 1','Heading 2'):continue
    title=normalize(p.text)
    hits=[i+1 for i,text in enumerate(pages) if title in text and 'СОДЕРЖАНИЕ' not in text]
    if not hits:raise RuntimeError(f'Heading not found in rendered PDF: {title}')
    entries.append((p,title,hits[0],int(p.style.name[-1])))
for name in ('TOC 1','TOC 2'):
    style=doc.styles[name] if name in doc.styles else doc.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
    style.font.name='Times New Roman';style.font.size=Pt(12)
    style.paragraph_format.first_line_indent=Cm(0)
    style.paragraph_format.left_indent=Cm(.5 if name=='TOC 2' else 0)
    style.paragraph_format.line_spacing=1.15
    style.paragraph_format.space_after=Pt(5)
body=doc.element.body
contents=next(p for p in doc.paragraphs if p.text=='СОДЕРЖАНИЕ')
start=body.index(contents._p)+1
end=body.index(entries[0][0]._p)
for el in list(body)[start:end]:body.remove(el)

def field(run,kind,instr=None):
    el=OxmlElement('w:instrText' if instr is not None else 'w:fldChar')
    if instr is not None:el.set(qn('xml:space'),'preserve');el.text=instr
    else:el.set(qn('w:fldCharType'),kind)
    run.append(el)

for i,(target,title,page,level) in enumerate(entries):
    bookmark=f'_TocDZ1_{i+1}'
    for el in list(target._p):
        if el.tag in (qn('w:bookmarkStart'),qn('w:bookmarkEnd')):target._p.remove(el)
    begin=OxmlElement('w:bookmarkStart');begin.set(qn('w:id'),str(100+i));begin.set(qn('w:name'),bookmark)
    finish=OxmlElement('w:bookmarkEnd');finish.set(qn('w:id'),str(100+i))
    target._p.insert(1 if target._p.pPr is not None else 0,begin);target._p.append(finish)
    p=doc.add_paragraph(style=f'TOC {level}')
    body.remove(p._p);body.insert(start+i,p._p)
    tabs=p.paragraph_format.tab_stops
    from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER
    tabs.add_tab_stop(Cm(15.5 if level==1 else 15),WD_TAB_ALIGNMENT.RIGHT,WD_TAB_LEADER.DOTS)
    if i==0:
        field(p.add_run()._r,'begin');field(p.add_run()._r,None,' TOC \\o "1-2" \\h \\z \\u ');field(p.add_run()._r,'separate')
    link=OxmlElement('w:hyperlink');link.set(qn('w:anchor'),bookmark)
    r=OxmlElement('w:r');t=OxmlElement('w:t');t.text=title;r.append(t);link.append(r)
    r=OxmlElement('w:r');r.append(OxmlElement('w:tab'));link.append(r)
    r=OxmlElement('w:r');field(r,'begin');link.append(r)
    r=OxmlElement('w:r');field(r,None,f' PAGEREF {bookmark} \\h ');link.append(r)
    r=OxmlElement('w:r');field(r,'separate');link.append(r)
    r=OxmlElement('w:r');t=OxmlElement('w:t');t.text=str(page);r.append(t);link.append(r)
    r=OxmlElement('w:r');field(r,'end');link.append(r)
    p._p.append(link)
    if i==len(entries)-1:field(p.add_run()._r,'end')
doc.save(path)
print('Generated live TOC field cache:',[(title,page) for _,title,page,_ in entries])
