"""Build a report from the original template and actual measurements."""
import json
from pathlib import Path
from statistics import mean
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
DATA=json.loads((ROOT/'results/measurements.json').read_text())
HOST=json.loads((ROOT/'results/host.json').read_text())
SMALL=DATA['datasets']['small']; WORK=DATA['datasets']['work']
OUT=ROOT/'Отчёт ДЗ1 — Невокшенов.docx'
doc=Document(next(ROOT.glob('Шаблон*.docx')))
body=doc.element.body
cover=[p._p for p in doc.paragraphs[:17]]
for el in list(body):
    if el not in cover and el.tag!=qn('w:sectPr'): body.remove(el)
doc.paragraphs[13].text='Вариант 3'
doc.paragraphs[14].text='Студент группы БИСТ-24-ПО-1\t_______________\tНевокшенов Е. А.'
section=doc.sections[0]
section.page_width=Cm(21);section.page_height=Cm(29.7)
section.top_margin=Cm(2);section.bottom_margin=Cm(2);section.left_margin=Cm(3);section.right_margin=Cm(1.5)
normal=doc.styles['Normal'];normal.font.name='Times New Roman';normal.font.size=Pt(14)
normal.paragraph_format.line_spacing=1.5
normal.paragraph_format.space_after=Pt(0)
normal.paragraph_format.first_line_indent=Cm(1.25)
for style in ('Heading 1','Heading 2'):
    s=doc.styles[style];s.font.name='Times New Roman';s.font.size=Pt(14);s.font.color.rgb=RGBColor(0,0,0)
    s.paragraph_format.first_line_indent=Cm(0);s.paragraph_format.space_before=Pt(14);s.paragraph_format.space_after=Pt(10)
    s.paragraph_format.keep_with_next=True
for p in doc.paragraphs:
    p.paragraph_format.first_line_indent=Cm(0)
# Page footer; cover is unnumbered.
section.different_first_page_header_footer=True
footer=section.footer.paragraphs[0];footer.clear();footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
footer.paragraph_format.first_line_indent=Cm(0)
r=footer.add_run();field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');r._r.addnext(field)

def paragraph(text='',style=None):
    p=doc.add_paragraph(text,style)
    if not style:p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    return p

def heading(text,level=1,new_page=False):
    p=doc.add_heading(text,level)
    if new_page:p.paragraph_format.page_break_before=True
    return p

def table(caption,headers,rows,widths=None):
    p=paragraph(caption);p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.keep_with_next=True;p.paragraph_format.space_before=Pt(6)
    t=doc.add_table(rows=1, cols=len(headers));t.style='Table Grid';t.autofit=False
    if widths:
        for col,width in zip(t.columns,widths):col.width=Cm(width)
    for cell,text in zip(t.rows[0].cells,headers):cell.text=str(text)
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        for c,text in zip(t.add_row().cells,row):c.text=str(text)
    for row in t.rows:
        no_split=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(no_split)
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.line_spacing=1.05;p.paragraph_format.space_after=Pt(4);p.paragraph_format.space_before=Pt(4)
                for r in p.runs:r.font.name='Times New Roman';r.font.size=Pt(10 if len(headers)>5 else 11)
    for c in t.rows[0].cells:
        for p in c.paragraphs:
            for r in p.runs:r.bold=True
    paragraph()
    return t

def code(text):
    for line in text.splitlines():
        p=paragraph();p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.line_spacing=1.0
        r=p.add_run(line);r.font.name='Courier New';r.font.size=Pt(9)

def picture(name,caption):
    p=paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Cm(0)
    p.paragraph_format.keep_with_next=True
    p.add_run().add_picture(str(ROOT/'results'/name),width=Cm(15.5))
    p=paragraph(caption);p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Cm(0);p.runs[0].font.size=Pt(12)

def fmt(x):return f'{x:.2f}'.replace('.',',')

NAMES={'login':'Вход','logout':'Выход','me':'Пользователь','services':'Услуги','specialists':'Специалисты','list':'Список записей','card':'Карточка','slots':'Свободные слоты','summary':'Сводка','create':'Бронирование','cancel':'Отмена'}
TASKS=['Реализовать сервис записи на приём с авторизацией и клиентским интерфейсом.','Описать контракт программного интерфейса.','Проверить бизнес-правила и операции интерфейса тестами.','Подготовить воспроизводимые малое и рабочее наполнения PostgreSQL.','Измерить время ответа всех операций на обоих объёмах.','Разложить серверное время операций между обращениями к БД и приложением.','Сформулировать гипотезы об узких местах по полученным числам.','Проверить очевидные ошибки реализации и зафиксировать результат проверки.']

p=paragraph('СОДЕРЖАНИЕ');p.paragraph_format.page_break_before=True;p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Cm(0)
r=p.add_run();r.bold=True
p=paragraph();p.paragraph_format.first_line_indent=Cm(0)
for kind,text in [('begin',None),(None,' TOC \\o "1-2" \\h \\z \\u '),('separate',None)]:
    r=p.add_run()
    if kind:
        e=OxmlElement('w:fldChar');e.set(qn('w:fldCharType'),kind)
    else:
        e=OxmlElement('w:instrText');e.set(qn('xml:space'),'preserve');e.text=text
    r._r.append(e)
p.add_run('Обновление автоматического содержания')
r=p.add_run();e=OxmlElement('w:fldChar');e.set(qn('w:fldCharType'),'end');r._r.append(e)
settings=doc.settings.element
update=OxmlElement('w:updateFields');update.set(qn('w:val'),'true');settings.append(update)

heading('ЗАДАНИЕ И ВАРИАНТ',new_page=True)
table('Таблица 1 – Вариант и стенд',['Параметр','Значение'],[
 ['Вариант и предметная область','3 — запись на приём к специалисту'],['Префикс','evgenii_nevokshenov'],['Язык и каркас','Python 3.12.10, FastAPI 0.115.12'],['Доступ к данным','psycopg 3.2.9, явные SQL-запросы'],['СУБД',DATA['metadata']['postgres'].split(',')[0]],['Клиентская часть','HTML, CSS, JavaScript; выдаётся сервером'],['Развёртывание','Docker Compose; PostgreSQL и приложение'],['Репозиторий','Публикация запланирована отдельно; ссылка добавляется перед сдачей.']], [5,10.5])
paragraph('Сервис хранит расписание специалистов и показывает свободные слоты для выбранной услуги. Бронирование отклоняется при пересечении с активной записью, неподходящей длительности или недоступности услуги у специалиста. Отмена сохраняет историю и освобождает время. Сводка показывает загрузку специалистов за выбранный период и долю отменённых записей.')

heading('ВВЕДЕНИЕ',new_page=True)
paragraph('Исходные показатели производительности нужны для последующего доказательства эффекта оптимизации. Оценка по субъективному впечатлению от интерфейса не позволяет отличить улучшение кода от изменения объёма данных, прогрева процесса или условий запуска. Поэтому до оптимизации фиксируются технологии, данные, сценарии и условия испытаний, а время измеряется серией повторов.')
paragraph('В отличие от лабораторного учебного стенда объектом исследования является собственный сервис записи на приём. Сначала подтверждается правильность его поведения, затем измеряется зависимость времени ответа от объёма и устанавливаются направления дальнейшего исследования.')
paragraph('Цель работы — разработать сервис по варианту 3 и получить воспроизводимые исходные показатели его производительности на двух объёмах данных.')
paragraph('Для достижения цели необходимо решить следующие задачи:')
for i,t in enumerate(TASKS,1): paragraph(f'{i}. {t}')

heading('1 Сервис, стенд и данные',new_page=True)
heading('1.1 Сервис и контракт',2)
paragraph('Реализованы вход, список записей с фильтрами и пагинацией, карточка, поиск слотов с бронированием и сводка. Авторизация использует серверную сессию в HttpOnly cookie, действующую восемь часов. Демо-пользователь demo выполняет все сценарии; пароль хранится в виде PBKDF2-хеша. Пересекающиеся бронирования одного специалиста сериализуются блокировкой строки специалиста в транзакции.')
paragraph('Основные сущности: специалист, услуга, слот и запись. Поддерживаемые услуги заданы массивом идентификаторов в специалисте. Пользователи и сессии являются техническими таблицами. При отмене запись получает статус cancelled, а последующее бронирование проверяет только активные записи.')
rows=[['POST /login','username, password','id, username; cookie','401, 422'],['POST /logout','cookie','ok=true','401'],['GET /me','cookie','id, username','401'],['GET /services','—','items: услуги','401'],['GET /specialists','—','items: специалисты','401'],['GET /appointments','page, size, status, specialist_id, from_at, to_at','items, total, page, size','401, 422'],['GET /appointments/{id}','id','карточка со связями','401, 404, 422'],['GET /slots','service_id, specialist_id, период, page, size','items, total, page, size','401, 404, 422'],['POST /appointments','slot_id, service_id, client_name, start_at?','201; карточка','401, 404, 409, 422'],['POST /appointments/{id}/cancel','id','карточка; cancelled','401, 404, 409, 422'],['GET /summary','from_at, to_at','показатели и специалисты','401, 422']]
table('Таблица 2 – Операции программного интерфейса',['Метод и путь /api','Параметры','Ответ','Ошибки'],rows,[4.2,4.5,4.6,2.2])
paragraph('Полный контракт с полями карточки и сводки находится в docs/API.md. Размер страницы составляет от 1 до 100 строк. Даты содержат часовой пояс; интервал периода полуоткрытый. Поиск слотов ограничен 31 днём. Отсутствие авторизации даёт 401, ошибки параметров — 422, отсутствующая запись — 404, конфликт бронирования или повторная отмена — 409. Кеширование не используется.')
picture('booking.jpg','Рисунок 1 – Поиск свободных слотов и переход к бронированию')
picture('summary.jpg','Рисунок 2 – Сводка на малом наборе после проверки интерфейса')
paragraph('Скриншоты получены при ручной проверке интерфейса. В основной базе дополнительно создана и отменена одна проверочная запись: на рисунке 2 поэтому 301 запись. Измерения выполнялись независимо, в отдельной базе evgenii_nevokshenov_measure.')

heading('1.2 Модульные тесты',2,new_page=True)
paragraph('Проверяются бизнес-правила интервалов и длительности, доступность услуги, расчёт долей, авторизация, поля ответов, фильтры, пагинация, отмена и конкурентное бронирование. Дополнительно проверены повторяемость наполнения и обработка ошибок измерителя. API-тесты используют реальную отдельную PostgreSQL, а чистые правила не обращаются к БД.')
paragraph('Команда запуска: bash scripts/control.sh tests. Итоговый прогон: 33 passed. Тестовая БД evgenii_nevokshenov_test изолирована от данных приложения.')
paragraph('Листинг 1 – Примеры тестов бизнес-правила и контракта')
code('''def test_adjacent_intervals_are_allowed():
    assert not overlaps(START, START+timedelta(hours=1),
                        START+timedelta(hours=1),
                        START+timedelta(hours=2))

def test_authorization_is_required(client):
    for path in ['/api/me', '/api/appointments',
                 '/api/services', '/api/specialists',
                 '/api/slots?service_id=1', '/api/summary']:
        assert client.get(path).status_code == 401''')
paragraph('Тест конкурентного бронирования выполняет два одновременных запроса и получает ровно один ответ 201 и один 409. Проверена также обрезка занятых минут границами периода; пустая сводка возвращает нулевые доли.')

heading('1.3 Условия испытания',2)
table('Таблица 3 – Условия испытания',['Что','Значение'],[
 ['Машина',f"{HOST['cpu']}; 18 ГБ RAM; macOS 26.6.2"],['Среда выполнения',f"Docker {HOST['docker']}, Linux ARM64; 11 CPU, около 7,65 ГиБ RAM для Docker; один процесс uvicorn"],['Учётная запись','demo; одинаковая для всех запросов'],['Повторы','30 измеряемых запросов для каждой операции и каждого объёма'],['Прогрев','5 успешных запросов перед каждой серией; исключены из расчётов'],['Инструмент','httpx 0.28.1, perf_counter; последовательный HTTP по loopback внутри контейнера'],['Периоды','Слоты: 1–8 октября 2026; список и сводка: 1 октября 2026 — 1 января 2028'],['Процентили','p50 — медиана; p95 — ближайший ранг ceil(0,95 × N); максимум по 30 запросам'],['Прочие процессы',HOST['other_containers']]], [4.5,11])
paragraph('Оба объёма исследовались одной версией кода на одном стенде. Замеры не являются нагрузочным испытанием: конкурентность равна единице, кроме отдельного теста корректности. Точные версии зависимостей записаны в requirements.lock.txt и results/measurements.json. Данные и показатели получены 4 октября 2026 года.')

heading('1.4 Наполнение базы данных',2,new_page=True)
paragraph('Наполнение выполняет команда python -m scripts.seed --size small|work --reset. Для обычного сервиса она запускается через docker compose exec app. Скрипт измерений вызывает тот же генератор в отдельной исследовательской БД и дважды формирует каждый объём; контрольные суммы всех четырёх таблиц совпали.')
table('Таблица 4 – Объём исходного наполнения',['Таблица','Малое','Рабочее','Рост'],[[t,SMALL['counts'][t],WORK['counts'][t],f"{WORK['counts'][t]/SMALL['counts'][t]:g} раз"] for t in ('specialists','services','slots','appointments')],[5.5,3,3.5,3.5])
paragraph('Генерация имеет фиксированную дату и порядок; активные записи не пересекаются. Исходная доля отменённых записей равна 20%. Фактические количества получены запросами count(*) к PostgreSQL, а контрольные суммы — из упорядоченного представления строк. Личные имена и сведения в базе синтетические.')
paragraph('Для измерения операций записи к каждому набору добавлено по 70 слотов и подготовлено по 35 записей для отмены. Прогрев создания добавил 5 записей, измеряемая серия — ещё 30; итого добавлено 70 записей, включая подготовку отмены. В конце малый набор содержит 570 слотов и 370 записей, рабочий — 200070 слотов и 60070 записей. Чтения измерены до серий создания и отмены, но после подготовки 35 записей отмены. Эти одинаковые добавки зафиксированы отдельно и не скрыты в объёме исходного наполнения.')

heading('2 Измерения',new_page=True)
heading('2.1 Время ответа операций',2)
paragraph('Таблица показывает полное время получения HTTP-ответа в миллисекундах. Все 11 операций контракта измерены на обоих объёмах. Исходные 30 значений каждой серии сохранены в results/measurements.json; значения прогрева в таблицу не включены.')
rows=[]
for key,w in WORK['operations'].items():
    a=SMALL['operations'][key]
    rows.append([NAMES[key],*[fmt(a['http_ms'][x]) for x in ('p50','p95','max')],*[fmt(w['http_ms'][x]) for x in ('p50','p95','max')],fmt(w['http_ms']['p50']/a['http_ms']['p50'])])
table('Таблица 5 – Время ответа на двух объёмах данных в мс',['Операция','p50 мал.','p95 мал.','max мал.','p50 раб.','p95 раб.','max раб.','Рост p50'],rows,[3.1,1.75,1.75,1.75,1.75,1.75,1.75,1.9])
ratios={k:WORK['operations'][k]['http_ms']['p50']/SMALL['operations'][k]['http_ms']['p50'] for k in NAMES}
paragraph(f"Список замедлился в {fmt(ratios['list'])} раза, свободные слоты — в {fmt(ratios['slots'])}, сводка — в {fmt(ratios['summary'])}. Это главные кандидаты на дальнейшее исследование. Объём записей вырос в 200 раз, слотов — в 400 раз; ни одна из этих операций не выросла быстрее исходных данных, однако их абсолютные затраты существенно увеличились.")
paragraph(f"Вход, выход, текущий пользователь, услуги, карточка и отмена остаются близкими по времени и исключены из дальнейшего разбора зависимости от объёма. Список специалистов вырос в {fmt(ratios['specialists'])} раза, создание — в {fmt(ratios['create'])}; они включены в разложение как второстепенные случаи. Время входа определяется прежде всего проверкой PBKDF2, а не числом записей на приём.")

heading('2.2 Разложение времени между приложением и базой данных',2)
paragraph('На рабочем объёме использованы серверные метрики Server-Timing. SQL-время измеряется вокруг execute и fetchall: включает ожидание PostgreSQL, передачу строк и их декодирование драйвером. Оно не является чистым временем исполнения внутри PostgreSQL. Время кода — остаток серверного времени после SQL; сюда входят открытие соединения, транзакционное завершение, бизнес-логика и подготовка ответа. Внешняя доставка HTTP измерена отдельно в таблице 5.')
paragraph('Для аддитивности разложения таблица 6 содержит арифметические средние по одним и тем же 30 запросам: серверное время равно сумме SQL и остатка. Медианы отдельных частей не складывались. Число SQL-запросов в каждой серии постоянно и включает авторизацию.')
rows=[]
for key in ('list','slots','summary','specialists','create'):
    samples=WORK['operations'][key]['samples']
    server=mean(s['server_ms'] for s in samples);sql=mean(s['sql_ms'] for s in samples);code_ms=mean(s['code_ms'] for s in samples)
    conclusion='БД' if sql>code_ms else 'Код' if key!='summary' else 'БД и код'
    if key=='summary':conclusion='БД и код'
    rows.append([NAMES[key],fmt(server),fmt(sql),fmt(code_ms),samples[0]['queries'],conclusion])
table('Таблица 6 – Куда уходит серверное время операции',['Операция','Всего, мс','SQL, мс','Код, мс','Запросов','Вывод'],rows,[3.5,2.3,2.3,2.3,2,3.1])
paragraph('У списка и поиска слотов доминируют обращения к БД. Сводка тратит значимое время на обе части: большие наборы строк получают из БД и затем обрабатывают в Python. У специалистов большая часть небольшого времени остаётся вне SQL, в том числе при сериализации ответа. Семь запросов бронирования имеют фиксированные назначения: авторизация, слот, блокировка специалиста, услуга, конфликт, вставка, карточка. Запросов в цикле по строкам ответа нет; число запросов не растёт с размером базы.')

heading('3 Гипотезы об узких местах',new_page=True)
for idx,key,cause in [
 (1,'list','COUNT(*) и сортировка списка без дополнительных индексов просматривают существенно больше строк на рабочем объёме. Фиксированные три SQL-запроса исключают объяснение ростом их количества. Проверка гипотезы в следующих заданиях: EXPLAIN (ANALYZE, BUFFERS) и отдельное исследование подсчёта и выборки страницы.'),
 (2,'slots','Поиск выполняет два соединения расписания с активными записями — для количества и страницы — и проверяет пересечения. При большом расписании и отсутствии дополнительных индексов стоимость фильтрации и соединения растёт. Проверка: планы обоих запросов и статистика прочитанных строк; затем индекс или изменение запроса с доказательством эффекта.'),
 (3,'summary','Приложение получает все подходящие слоты и записи, объединяет интервалы и агрегирует показатели в Python. Существенны и передача/декодирование строк, и обработка больших коллекций. Проверка: профиль Python, измерение размера ответа БД и сравнение с агрегацией на стороне PostgreSQL без изменения формул сводки.')]:
    a=SMALL['operations'][key];w=WORK['operations'][key]
    paragraph(f'Гипотеза {idx}. {NAMES[key]}.')
    paragraph(f"1. Операция: {w['method']} {w['path'].split('?')[0]}.")
    paragraph(f"2. Основание: HTTP p50 {fmt(a['http_ms']['p50'])} → {fmt(w['http_ms']['p50'])} мс; p95 на рабочем объёме {fmt(w['http_ms']['p95'])} мс. Медиана серверного времени {fmt(w['server_ms']['p50'])} мс, SQL {fmt(w['sql_ms']['p50'])} мс, остаток {fmt(w['code_ms']['p50'])} мс; {int(w['queries']['p50'])} запроса на операцию.")
    paragraph('3. Предполагаемая причина: '+cause)
paragraph('При проверке исходной схемы обнаружен дополнительный индекс username, автоматически созданный ограничением UNIQUE. Он удалён, поскольку в ДЗ1 допустимы только индексы ключей; отдельный тест проверяет отсутствие вторичных индексов. После изменения заново сняты все серии на обоих объёмах. Данные до исправления сохранены отдельно в results/measurements-before-index-fix.json. Изменение касается технической таблицы одного демо-пользователя, а не предметных запросов. Медиана входа на рабочем объёме до исправления составляла 27.24 мс, после — 26.53 мс; это различие в повторных сериях не доказывает ускорения. Все итоговые таблицы используют только повторный прогон.')
paragraph('Очевидные ошибки. В итоговой версии ошибок контракта, пропущенной пагинации или повторного одинакового SQL-запроса не обнаружено. Это подтверждено 33 тестами, ручным выполнением входа, поиска, записи, просмотра и отмены, а также постоянным числом SQL-запросов. Подсчёт общего количества и выборка страницы являются разными запросами, а не случайным дублем. Все представленные серии относятся к одной проверенной версии бизнес-логики. Сохраняющиеся затраты являются предметом дальнейшего профилирования и оптимизации, а не основанием намеренно ухудшать код.')

heading('ЗАКЛЮЧЕНИЕ',new_page=True)
paragraph('Цель работы достигнута: реализован сервис записи на приём по варианту 3 и получены исходные показатели на двух воспроизводимых объёмах данных.')
paragraph('Для этого решены следующие задачи:')
completed=['Реализован сервис с авторизацией, списком, карточкой, поиском слотов, бронированием, отменой и сводкой.','Описан контракт 11 операций программного интерфейса.','Выполнено 33 успешных тестов, включая конкурентное бронирование.','Подготовлены наборы с 300 и 60000 исходных записей, 500 и 200000 слотов; подтверждены повторные контрольные суммы.','Измерены 11 операций на двух объёмах: по 30 запросов после пяти прогревочных повторов, рассчитаны p50, p95 и максимум.','Разложено серверное время пяти операций; зафиксированы SQL-время, остаток и количество запросов.','Сформулированы три гипотезы по результатам измерений.','Проверены очевидные ошибки; нарушений контракта и лишних запросов в итоговой версии не выявлено.']
for i,text in enumerate(completed,1):paragraph(f'{i}. {text}')
paragraph(f"Основные затраты обнаружены в списке, поиске свободных слотов и сводке. На рабочем объёме их HTTP-медианы составили соответственно {fmt(WORK['operations']['list']['http_ms']['p50'])}, {fmt(WORK['operations']['slots']['http_ms']['p50'])} и {fmt(WORK['operations']['summary']['http_ms']['p50'])} мс. Для списка и слотов преобладают обращения к БД; для сводки значимы также обработка и агрегация в Python. Результаты дают исходную точку для нагрузочного исследования и последующей оптимизации при сохранении контракта.")
doc.save(OUT)
print(OUT)
